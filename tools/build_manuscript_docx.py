#!/usr/bin/env python3
"""Build a manuscript-formatted .docx from a pandoc Markdown file.

Formatting follows a conventional Nature Communications submission:
Arial 11 pt, 1.5 line spacing, black sentence-case headings, 2.54 cm margins,
continuous line numbers, centred page numbers, three-rule tables and page
breaks before References, Figure legends and Supplementary Table.

Usage: python3 tools/build_manuscript_docx.py <input.md> <output.docx>
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PANDOC = shutil.which("pandoc") or "pandoc"   # pandoc >= 3 on PATH (env/environment.yml)
FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
NEW_PAGE_HEADINGS = {"References", "Figure legends"}


def set_font(style_or_run, size=None, bold=None, italic=None):
    f = style_or_run.font
    f.name = FONT
    f.color.rgb = BLACK
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    # replace theme-font bindings so Arial applies to every script
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in list(rfonts.attrib):
        del rfonts.attrib[att]
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for att in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if color.get(qn(att)) is not None:
                del color.attrib[qn(att)]


def make_reference_docx(path):
    subprocess.run([str(PANDOC), "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(path)
    st = d.styles
    body = dict(size=11)
    for name in ("Normal", "Body Text", "First Paragraph", "Block Text"):
        if name in st:
            set_font(st[name], **body)
            pf = st[name].paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.space_before = Pt(0)
            pf.space_after = Pt(8)
            pf.first_line_indent = Cm(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if "Compact" in st:  # tight lists (references) and table cells
        set_font(st["Compact"], size=10)
        pf = st["Compact"].paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(4)
    set_font(st["Title"], size=16, bold=True)
    st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st["Title"].paragraph_format.space_after = Pt(6)
    set_font(st["Subtitle"], size=11, bold=False, italic=False)
    st["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st["Subtitle"].paragraph_format.space_after = Pt(18)
    for name, size in (("Heading 1", 13), ("Heading 2", 11), ("Heading 3", 11)):
        set_font(st[name], size=size, bold=True, italic=False)
        pf = st[name].paragraph_format
        pf.space_before = Pt(18 if name == "Heading 1" else 12)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for name in ("Hyperlink",):
        if name in st:
            set_font(st[name])
    d.save(path)


def three_rule_table(tbl):
    """Top rule, rule under the header row, bottom rule; no vertical lines."""
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        on = edge in ("top", "bottom")
        el.set(qn("w:val"), "single" if on else "nil")
        if on:
            el.set(qn("w:sz"), "8")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
    for cell in tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "000000")
        tcb.append(b)
        tcPr.append(tcb)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = FONT


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)
    run.font.name = FONT
    run.font.size = Pt(10)


def add_line_numbers(section):
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:distance"), "283")  # 0.5 cm
    ln.set(qn("w:restart"), "continuous")
    pgMar = sectPr.find(qn("w:pgMar"))
    pgMar.addnext(ln)  # schema order: pgSz, pgMar, ..., lnNumType, pgNumType, cols


def main(md, out):
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "reference.docx"
        make_reference_docx(ref)
        subprocess.run([str(PANDOC), str(md), "-o", str(out), "--reference-doc", str(ref)], check=True)
    d = Document(out)
    for s in d.sections:
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)  # US Letter, matching the main manuscript
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
        add_line_numbers(s)
        add_page_number_footer(s)
    in_refs = False
    for p in d.paragraphs:
        txt = p.text.strip()
        if p.style.name == "Heading 1":
            in_refs = txt == "References"
            if txt in NEW_PAGE_HEADINGS:
                p.paragraph_format.page_break_before = True
        elif in_refs and txt[:1].isdigit():  # numbered reference: 10 pt, single spaced, hanging indent
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(4)
            pf.left_indent = Cm(0.75)
            pf.first_line_indent = Cm(-0.75)
            for r in p.runs:
                r.font.size = Pt(10)
        elif txt.startswith("Supplementary Table") and "|" in txt:  # table caption opens its own page
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
    for tbl in d.tables:
        three_rule_table(tbl)
        tblW = tbl._tbl.tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tbl._tbl.tblPr.append(tblW)
        tblW.set(qn("w:type"), "pct")
        tblW.set(qn("w:w"), "5000")  # full text width
        # rescale pandoc's grid so the columns fill the text block, keeping their proportions
        text_w = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin
        grid = tbl._tbl.tblGrid.findall(qn("w:gridCol"))
        old = [int(g.get(qn("w:w"))) for g in grid]
        new = [int(text_w * w / sum(old)) for w in old]
        for g, w in zip(grid, new):
            g.set(qn("w:w"), str(int(w / 635)))  # EMU -> twips
        for row in tbl.rows:
            for cell, w in zip(row.cells, new):
                cell.width = w
    d.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
